"""把 PDF 重建成可以編輯的 Word 檔。

PDF 只記「哪個字畫在哪個座標」,沒有段落、表格的概念,所以這裡是「重建」而不是「翻譯」:

1. 取出每個字的位置、字級、字型、顏色
2. 用文字底線合併成行(用方框會把「一」「，」這種又扁又低的字丟到別行)
3. 依行距、左緣、字級合併成段落
4. 上下位置重疊、左右分開的內容歸成同一列;並排的圖片、分欄、表格都靠這一步
5. 連續幾列而且欄位對得起來的,輸出成 Word 表格(原檔有框線才畫框線)
6. 用 python-docx 寫出段落、字型、顏色、圖片,每一頁一個版面設定

這樣產生的是真正的段落與表格,可以直接接著編輯;用 LibreOffice 轉則會變成一大堆文字方塊。
"""

import ctypes
import io
import statistics
from dataclasses import dataclass, field
from pathlib import Path

import pypdfium2.raw as raw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor

from core import pdfium
from core.files import free_path

EMU_PER_POINT = 12700
SCAN_DPI = 150              # 沒有文字的頁面(掃描檔)直接放整頁圖片時用的解析度
MIN_MARGIN = 14.0           # 邊界至少留這麼多點,避免文字貼著紙邊
COLUMN_TOLERANCE = 8.0      # 左邊界差這麼多點以內就當成同一欄
OCR_SPLIT_GAP = 1.5         # 辨識出來的詞之間空到這麼多個字寬,就當成表格或分欄的不同格
VECTOR_DPI = 200            # 圖表(用線條畫出來的)整塊截圖時的解析度
VECTOR_GAP = 8.0            # 線條相距這麼多點以內就當成同一張圖
VECTOR_MIN_PATHS = 8        # 這麼多條線以上才當成圖表,不然只是表格框線
VECTOR_MIN_SIZE = 40.0      # 太小的線條群不截圖
CJK_RANGES = ((0x2E80, 0xA4CF), (0xF900, 0xFAFF), (0xFF00, 0xFF60), (0x20000, 0x3FFFF))
# PDF 裡的中文字型常常沒有名字,只好依「有沒有襯線」挑一個系統字型
CJK_FALLBACK = {"serif": "新細明體", "sans": "微軟正黑體"}
# 常見字型的 PostScript 名稱 → Word 看得懂的名稱
FONT_NAMES = {
    "ArialMT": "Arial", "Arial-BoldMT": "Arial", "Arial-ItalicMT": "Arial",
    "TimesNewRomanPSMT": "Times New Roman", "TimesNewRomanPS-BoldMT": "Times New Roman",
    "TimesNewRomanPS-ItalicMT": "Times New Roman", "CourierNewPSMT": "Courier New",
    "Helvetica": "Arial", "Helvetica-Bold": "Arial", "SymbolMT": "Symbol",
}
_FLAG_SERIF = 1 << 1
_FLAG_ITALIC = 1 << 6
_FLAG_BOLD = 1 << 18


def is_cjk(ch) -> bool:
    code = ord(ch)
    return any(low <= code <= high for low, high in CJK_RANGES)


def clean_font(name: str) -> str:
    """去掉子集字型的前綴(ABCDEF+)與樣式後綴,換成 Word 慣用的名稱。"""
    if not name:
        return ""
    if len(name) > 7 and name[6] == "+":
        name = name[7:]
    if name in FONT_NAMES:
        return FONT_NAMES[name]
    for mark in ("-Bold", "-Italic", "-Oblique", "-BoldItalic", ",Bold", ",Italic", "-Regular", "MT", "PS"):
        if name.endswith(mark):
            name = name[: -len(mark)]
    return name.replace("-", " ").strip()


@dataclass
class Char:
    text: str
    left: float
    top: float
    right: float
    bottom: float
    baseline: float             # 文字底線;同一行的字底線一樣高,用這個分行最準
    size: float
    font: str
    bold: bool
    italic: bool
    color: tuple
    generated: bool = False     # PDFium 依字距自己補上的空白,不是檔案裡真的有的字

    @property
    def style(self):
        return round(self.size, 1), self.font, self.bold, self.italic, self.color


@dataclass
class Line:
    chars: list = field(default_factory=list)

    @property
    def baseline(self):
        return statistics.median([c.baseline for c in self.chars])

    @property
    def left(self):
        return min(c.left for c in self.chars)

    @property
    def right(self):
        return max(c.right for c in self.chars)

    @property
    def top(self):
        return min(c.top for c in self.chars)

    @property
    def bottom(self):
        return max(c.bottom for c in self.chars)

    @property
    def size(self):
        return statistics.median([c.size for c in self.chars])

    @property
    def text(self):
        return "".join(c.text for c in self.chars)


@dataclass
class Block:
    """一個段落或一張圖片。"""
    kind: str                       # text 或 image
    top: float
    left: float = 0.0
    right: float = 0.0
    lines: list = field(default_factory=list)
    image: bytes = b""
    width: float = 0.0
    height: float = 0.0


# ---------------------------------------------------------------- 讀取 PDF

def _read_chars(text_page, height) -> list:
    """取出頁面上每個字的位置與樣式;座標換成左上角為原點、往下為正。"""
    chars = []
    count = raw.FPDFText_CountChars(text_page)
    name_buffer = ctypes.create_string_buffer(128)
    for index in range(count):
        code = raw.FPDFText_GetUnicode(text_page, index)
        if code in (0, 10, 13):
            continue
        generated = code == 32 and raw.FPDFText_IsGenerated(text_page, index) == 1
        left, right, bottom, top = (ctypes.c_double() for _ in range(4))
        raw.FPDFText_GetCharBox(text_page, index, *[ctypes.byref(v) for v in (left, right, bottom, top)])
        flags = ctypes.c_int()
        raw.FPDFText_GetFontInfo(text_page, index, name_buffer, 128, ctypes.byref(flags))
        name = name_buffer.value.decode("utf-8", "replace")
        font = clean_font(name) or CJK_FALLBACK["serif" if flags.value & _FLAG_SERIF else "sans"]
        red, green, blue, alpha = (ctypes.c_uint() for _ in range(4))
        raw.FPDFText_GetFillColor(text_page, index, *[ctypes.byref(v) for v in (red, green, blue, alpha)])
        origin_x, origin_y = ctypes.c_double(), ctypes.c_double()
        raw.FPDFText_GetCharOrigin(text_page, index, ctypes.byref(origin_x), ctypes.byref(origin_y))
        chars.append(Char(chr(code), left.value, height - top.value, right.value, height - bottom.value,
                          height - origin_y.value, pdfium.char_size(text_page, index), font,
                          bool(flags.value & _FLAG_BOLD) or "Bold" in name,
                          bool(flags.value & _FLAG_ITALIC) or "Italic" in name or "Oblique" in name,
                          (red.value, green.value, blue.value), generated))
    return chars


def _drop_extra_spaces(chars) -> list:
    """PDFium 會依字距自己補空白:中文之間不需要,英文單字之間要留著。"""
    kept = []
    for index, char in enumerate(chars):
        if char.generated:
            before = chars[index - 1] if index else None
            after = chars[index + 1] if index + 1 < len(chars) else None
            if (before and is_cjk(before.text)) or (after and is_cjk(after.text)):
                continue
            # 字和字之間只是排版上的微小間隙時也不算空白,不然 PTR 會變成 P TR
            if before and after and after.left - before.right < max(before.size, after.size, 1.0) * 0.22:
                continue
        kept.append(char)
    return kept


def _group_lines(chars, split_gap=2.5) -> list:
    """把字合併成行。

    用文字底線判斷,不用方框:像「一」「，」這種又扁又低的字,方框和整行重疊很少,
    用方框判斷會被丟到別行去,整句就亂掉了。
    """
    lines = []
    for char in sorted(chars, key=lambda c: (round(c.baseline, 1), c.left)):
        if lines and abs(char.baseline - lines[-1].baseline) <= max(1.5, 0.4 * max(char.size, lines[-1].size)):
            lines[-1].chars.append(char)
        else:
            lines.append(Line([char]))
    result = []
    for line in lines:
        line.chars.sort(key=lambda c: c.left)
        line.chars = _drop_extra_spaces(line.chars)
        if not line.chars:
            continue
        # 同一條底線上中間空一大段,多半是分欄或表格的不同格子,拆開才不會讀錯順序
        current = [line.chars[0]]
        for previous, char in zip(line.chars, line.chars[1:]):
            if char.left - previous.right > max(previous.size, char.size) * split_gap:
                result.append(Line(current))
                current = []
            current.append(char)
        result.append(Line(current))
    return sorted(result, key=lambda l: (round(l.baseline, 1), l.left))


def _group_paragraphs(lines) -> list:
    """把行合併成段落:行距沒有明顯變大、左邊對齊差不多、字級接近、左右有重疊,就算同一段。"""
    blocks = []
    for line in lines:
        if blocks and blocks[-1].kind == "text":
            previous = blocks[-1].lines[-1]
            gap = line.top - previous.bottom
            height = max(previous.size, line.size)
            same_size = abs(previous.size - line.size) <= max(1.0, height * 0.25)
            aligned = abs(line.left - previous.left) <= height * 2.5 or line.left > previous.left
            overlap = min(blocks[-1].right, line.right) - max(blocks[-1].left, line.left)
            if gap <= height * 0.9 and same_size and aligned and overlap > 0:
                blocks[-1].lines.append(line)
                blocks[-1].left = min(blocks[-1].left, line.left)
                blocks[-1].right = max(blocks[-1].right, line.right)
                continue
        blocks.append(Block("text", line.top, line.left, line.right, [line]))
    return blocks


def _read_images(document, page, page_height) -> list:
    """頁面上的圖片:位置與內容(畫成 PNG);座標換成左上角為原點。"""
    from PIL import Image

    results = []
    for index in range(raw.FPDFPage_CountObjects(page.raw)):
        obj = raw.FPDFPage_GetObject(page.raw, index)
        if raw.FPDFPageObj_GetType(obj) != raw.FPDF_PAGEOBJ_IMAGE:
            continue
        left, bottom, right, top = (ctypes.c_float() for _ in range(4))
        raw.FPDFPageObj_GetBounds(obj, *[ctypes.byref(v) for v in (left, bottom, right, top)])
        width, box_height = right.value - left.value, top.value - bottom.value
        if width < 8 or box_height < 8:
            continue                # 太小的圖通常是線條或裝飾,跳過
        bitmap = raw.FPDFImageObj_GetRenderedBitmap(document.raw, page.raw, obj)
        if not bitmap:
            continue
        try:
            stride = raw.FPDFBitmap_GetStride(bitmap)
            size = (raw.FPDFBitmap_GetWidth(bitmap), raw.FPDFBitmap_GetHeight(bitmap))
            data = ctypes.string_at(raw.FPDFBitmap_GetBuffer(bitmap), stride * size[1])
            image = Image.frombuffer("RGBA", size, data, "raw", "BGRA", stride, 1).convert("RGB")
        finally:
            raw.FPDFBitmap_Destroy(bitmap)
        buffer = io.BytesIO()
        image.save(buffer, "PNG")
        results.append(Block("image", page_height - top.value, left.value, right.value,
                             image=buffer.getvalue(), width=width, height=box_height))
    return results


def _vector_boxes(page, page_width, page_height) -> list:
    """線條類物件的範圍(PDF 座標:左、下、右、上);背景色塊與整頁大的方框不算。"""
    boxes = []
    page_area = page_width * page_height
    for index in range(raw.FPDFPage_CountObjects(page.raw)):
        obj = raw.FPDFPage_GetObject(page.raw, index)
        if raw.FPDFPageObj_GetType(obj) != raw.FPDF_PAGEOBJ_PATH:
            continue
        left, bottom, right, top = (ctypes.c_float() for _ in range(4))
        raw.FPDFPageObj_GetBounds(obj, *[ctypes.byref(v) for v in (left, bottom, right, top)])
        box = (left.value, bottom.value, right.value, top.value)
        if (box[2] - box[0]) * (box[3] - box[1]) > page_area * 0.8:
            continue                # 整頁的底色或外框,不是圖表
        boxes.append(box)
    return boxes


def _cluster_boxes(boxes, gap) -> list:
    """把靠在一起的線條併成一塊;回傳 [(範圍, 幾個線條物件)]。"""
    clusters = []                   # [[left, bottom, right, top, count]]
    for box in sorted(boxes, key=lambda b: (-b[3], b[0])):
        for cluster in clusters:
            if (box[0] <= cluster[2] + gap and box[2] >= cluster[0] - gap
                    and box[1] <= cluster[3] + gap and box[3] >= cluster[1] - gap):
                cluster[0] = min(cluster[0], box[0])
                cluster[1] = min(cluster[1], box[1])
                cluster[2] = max(cluster[2], box[2])
                cluster[3] = max(cluster[3], box[3])
                cluster[4] += 1
                break
        else:
            clusters.append([*box, 1])
    merged = True
    while merged:                   # 前面併過之後範圍變大,可能又和別塊接上了
        merged = False
        for i in range(len(clusters)):
            for j in range(len(clusters) - 1, i, -1):
                a, b = clusters[i], clusters[j]
                if (b[0] <= a[2] + gap and b[2] >= a[0] - gap
                        and b[1] <= a[3] + gap and b[3] >= a[1] - gap):
                    a[0], a[1] = min(a[0], b[0]), min(a[1], b[1])
                    a[2], a[3] = max(a[2], b[2]), max(a[3], b[3])
                    a[4] += b[4]
                    clusters.pop(j)
                    merged = True
    return [((c[0], c[1], c[2], c[3]), c[4]) for c in clusters]


def _vector_pictures(page, chars, page_width, page_height) -> tuple:
    """把圖表(一大堆線條畫出來的)整塊截成圖片。

    PDF 裡的折線圖、長條圖是幾百個線條物件,文字標籤則散在圖上,而且常常上下重疊。
    照一般段落輸出會讓標籤擠進內文、把頁面撐開,所以整塊畫成圖片,裡面的文字就跟著不輸出。
    回傳 (圖片區塊, 留下來的字)。
    """
    clusters = _cluster_boxes(_vector_boxes(page, page_width, page_height), VECTOR_GAP)
    picked = []
    for (left, bottom, right, top), count in clusters:
        if count < VECTOR_MIN_PATHS or right - left < VECTOR_MIN_SIZE or top - bottom < VECTOR_MIN_SIZE:
            continue
        if (right - left) * (top - bottom) > page_width * page_height * 0.85:
            continue                # 幾乎整頁大,多半是外框或底紋,不是圖表
        inside = [c for c in chars
                  if left - 2 <= (c.left + c.right) / 2 <= right + 2
                  and bottom - 2 <= page_height - (c.top + c.bottom) / 2 <= top + 2]
        if len(inside) > len(chars) * 0.8:
            continue                # 整頁的字都在裡面,那是版面框線而不是圖表
        picked.append(((left, bottom, right, top), set(id(c) for c in inside)))
    if not picked:
        return [], chars

    blocks, removed = [], set()
    for (left, bottom, right, top), inside in picked:
        crop = (max(0.0, left - 2), max(0.0, bottom - 2),
                max(0.0, page_width - right - 2), max(0.0, page_height - top - 2))
        bitmap = page.render(scale=VECTOR_DPI / 72, crop=crop)
        buffer = io.BytesIO()
        bitmap.to_pil().convert("RGB").save(buffer, "PNG")
        bitmap.close()
        width = page_width - crop[0] - crop[2]
        height = page_height - crop[1] - crop[3]
        blocks.append(Block("image", page_height - top - 2, crop[0], crop[0] + width,
                            image=buffer.getvalue(), width=width, height=height))
        removed |= inside
    return blocks, [c for c in chars if id(c) not in removed]


def _read_rulings(page, page_height) -> list:
    """頁面上的框線:又細又長的線條;用來判斷表格要不要畫框線。"""
    rulings = []
    for index in range(raw.FPDFPage_CountObjects(page.raw)):
        obj = raw.FPDFPage_GetObject(page.raw, index)
        if raw.FPDFPageObj_GetType(obj) != raw.FPDF_PAGEOBJ_PATH:
            continue
        left, bottom, right, top = (ctypes.c_float() for _ in range(4))
        raw.FPDFPageObj_GetBounds(obj, *[ctypes.byref(v) for v in (left, bottom, right, top)])
        width, thickness = right.value - left.value, top.value - bottom.value
        if (thickness <= 2.5 and width >= 24) or (width <= 2.5 and thickness >= 24):
            rulings.append((left.value, page_height - top.value, right.value, page_height - bottom.value))
    return rulings


# ---------------------------------------------------------------- 版面

def _block_span(block):
    """一個區塊在頁面上的上下範圍。"""
    if block.kind == "image":
        return block.top, block.top + block.height
    return block.lines[0].top, block.lines[-1].bottom


def _group_rows(blocks) -> list:
    """把上下重疊、左右分開的內容歸成同一列。

    投影片並排的圖片、分欄的文字、表格的一列都屬於這種;少了這一步,並排的東西會被疊成上下,
    一頁塞不下就會多跑出一頁。
    """
    rows = []
    for block in sorted(blocks, key=lambda b: (round(_block_span(b)[0], 1), b.left)):
        top, bottom = _block_span(block)
        if rows:
            row = rows[-1]
            overlap = min(bottom, row["bottom"]) - max(top, row["top"])
            smallest = min(bottom - top, row["bottom"] - row["top"]) or 1.0
            apart = all(block.left >= other.right - 1 or block.right <= other.left + 1 for other in row["blocks"])
            if overlap > smallest * 0.5 and apart:
                row["blocks"].append(block)
                row["top"], row["bottom"] = min(row["top"], top), max(row["bottom"], bottom)
                continue
        rows.append({"blocks": [block], "top": top, "bottom": bottom})
    for row in rows:
        row["blocks"].sort(key=lambda b: b.left)
    return rows


def _same_columns(first, second) -> bool:
    if len(first["blocks"]) != len(second["blocks"]):
        return False
    return all(abs(a.left - b.left) <= COLUMN_TOLERANCE for a, b in zip(first["blocks"], second["blocks"]))


def _group_tables(rows) -> list:
    """連續幾列而且欄位對得起來的,合併成一個表格;其他維持單獨一列。"""
    groups, current = [], []
    for row in rows:
        if len(row["blocks"]) < 2:
            if current:
                groups.append(current)
                current = []
            groups.append([row])
            continue
        if current and _same_columns(current[-1], row):
            current.append(row)
        else:
            if current:
                groups.append(current)
            current = [row]
    if current:
        groups.append(current)
    return groups


def _table_columns(group, content_right):
    """表格每一欄的左邊界與寬度。"""
    lefts = sorted({round(block.left, 1) for row in group for block in row["blocks"]})
    grouped = []
    for left in lefts:
        if grouped and left - grouped[-1][-1] <= COLUMN_TOLERANCE:
            grouped[-1].append(left)
        else:
            grouped.append([left])
    edges = [min(item) for item in grouped]
    widths = [(edges[i + 1] if i + 1 < len(edges) else content_right) - edges[i] for i in range(len(edges))]
    return edges, [max(width, 12.0) for width in widths]


def _column_index(edges, block) -> int:
    return min(range(len(edges)), key=lambda index: abs(block.left - edges[index]))


def _has_rulings(group, rulings) -> bool:
    """這個表格範圍裡有沒有原本就畫好的框線。"""
    if not rulings:
        return False
    top = min(row["top"] for row in group) - 8
    bottom = max(row["bottom"] for row in group) + 8
    left = min(block.left for row in group for block in row["blocks"]) - 8
    right = max(block.right for row in group for block in row["blocks"]) + 8
    inside = [r for r in rulings if r[0] >= left and r[2] <= right and r[1] >= top and r[3] <= bottom]
    return len(inside) >= 2


# ---------------------------------------------------------------- 寫成 Word

def _prepare_document(word):
    """Word 預設樣式本身有段後間距和 1.08 行距,會把內容撐高、整份多出好幾頁。
    這裡全部歸零,行距改由每一段依原檔實際量到的行高決定。"""
    normal = word.styles["Normal"].paragraph_format
    normal.space_before = Pt(0)
    normal.space_after = Pt(0)
    normal.line_spacing = 1.0
    normal.widow_control = False      # 不要為了避免單行落單自己換頁


def _line_pitch(block, page_pitch=0.0) -> float:
    """段落裡相鄰兩行底線的距離。

    只有一行時沒有行距可量,改用整頁的行距推估:直接用「字級 × 1.2」會比原檔大一點點,
    每一段多幾分點,整頁累積下來就會多擠出一頁。
    """
    baselines = [line.baseline for line in block.lines]
    if len(baselines) >= 2:
        gaps = [b - a for a, b in zip(baselines, baselines[1:]) if b > a]
        if gaps:
            return statistics.median(gaps)
    size = block.lines[0].size
    if page_pitch and size * 0.9 <= page_pitch <= size * 2.0:
        return page_pitch           # 整頁的行距和這一段的字級對得起來才採用
    return max(size * 1.2, 1.0)


def _page_pitch(blocks) -> float:
    """整頁相鄰兩行底線的距離(中位數);拿來當單行段落的行高。"""
    baselines = sorted(line.baseline for block in blocks if block.kind == "text" for line in block.lines)
    gaps = [b - a for a, b in zip(baselines, baselines[1:]) if 1.0 < b - a < 60]
    return statistics.median(gaps) if gaps else 0.0


def _apply_style(run, char_style, default_size):
    size, font, bold, italic, color = char_style
    if font:
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)      # 中文字也要指定,不然會用預設字型
    run.font.size = Pt(size if size > 1 else default_size)
    run.bold = bold
    run.italic = italic
    if color != (0, 0, 0):
        run.font.color.rgb = RGBColor(*color)


def _paragraph_text(block) -> str:
    pieces = []
    for index, line in enumerate(block.lines):
        text = line.text.strip()
        if index and pieces:
            previous = pieces[-1][-1] if pieces[-1] else ""
            if previous and text and not (is_cjk(previous) or is_cjk(text[0])):
                pieces.append(" ")
        pieces.append(text)
    return "".join(pieces)


def _alignment(block, content_left, content_right):
    """依位置判斷置中、靠右,其他都當靠左。"""
    width = content_right - content_left
    if width <= 0:
        return None
    left_gap, right_gap = block.left - content_left, content_right - block.right
    if left_gap > width * 0.08 and abs(left_gap - right_gap) < width * 0.06:
        return WD_ALIGN_PARAGRAPH.CENTER
    if left_gap > width * 0.25 and right_gap < width * 0.03:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return None


def _clear_cell_padding(table):
    """把表格的內距設成 0,內容才會停在原本的位置。"""
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    table._tbl.tblPr.append(margins)


def _write_block(container, block, layout, state):
    """把一個段落或圖片寫進文件,或寫進表格的儲存格。"""
    if block.kind == "image":
        paragraph = container.add_paragraph()
        paragraph.alignment = _alignment(block, layout["left"], layout["right"])
        spacing = paragraph.paragraph_format
        spacing.space_after = Pt(0)
        if state["baseline"] is not None:
            spacing.space_before = Pt(max(0.0, min(48.0, round(block.top - state["baseline"], 1))))
        paragraph.add_run().add_picture(io.BytesIO(block.image), width=Emu(round(block.width * EMU_PER_POINT)))
        state["baseline"] = block.top + block.height
        return
    if not _paragraph_text(block).strip():
        return
    paragraph = container.add_paragraph()
    paragraph.alignment = _alignment(block, layout["left"], layout["right"])
    pitch = _line_pitch(block, layout.get("pitch", 0.0))
    spacing = paragraph.paragraph_format
    spacing.space_after = Pt(0)
    spacing.line_spacing = Pt(round(pitch, 1))                  # 行距照原檔
    if layout["slack"]:
        # 換用的字型通常比原檔稍寬,留一點右側餘裕,才不會為了差幾個字就多折一行
        spacing.right_indent = Pt(-layout["slack"])
    if state["baseline"] is not None:
        # Word 本來就會空一行的高度,只要再補上多出來的距離,不然每段都會多佔一行
        extra = block.lines[0].baseline - state["baseline"] - pitch
        spacing.space_before = Pt(max(0.0, min(48.0, round(extra, 1))))
    _write_runs(paragraph, block)
    state["baseline"] = block.lines[-1].baseline


def _write_runs(paragraph, block, breaks=False):
    """把段落的字依樣式分段寫進去;breaks 為真時照原檔換行(照原樣模式),否則讓 Word 自己折行。"""
    default_size = block.lines[0].size
    for index, line in enumerate(block.lines):
        if index and line.chars:
            if breaks:
                paragraph.add_run().add_break()
            else:
                previous = block.lines[index - 1].text.strip()[-1:]
                if previous and not (is_cjk(previous) or is_cjk(line.text.strip()[:1] or " ")):
                    paragraph.add_run(" ")
        buffer, style = "", None
        for char in line.chars:
            if style is None or char.style == style:
                buffer += char.text
                style = char.style
                continue
            _apply_style(paragraph.add_run(buffer), style, default_size)
            buffer, style = char.text, char.style
        if buffer and style is not None:
            _apply_style(paragraph.add_run(buffer), style, default_size)


def _write_table(word, group, layout, state, bordered):
    """一列裡有好幾塊內容(並排的圖片、分欄、表格)時,用表格排出來。"""
    edges, widths = _table_columns(group, layout["right"])
    table = word.add_table(rows=len(group), cols=len(edges))
    table.style = word.styles["Table Grid" if bordered else "Normal Table"]
    table.autofit = False
    _clear_cell_padding(table)
    for index, width in enumerate(widths):
        for cell in table.columns[index].cells:
            cell.width = Emu(round(width * EMU_PER_POINT))
    previous = state["baseline"]
    for row_index, row in enumerate(group):
        for block in row["blocks"]:
            cell = table.cell(row_index, _column_index(edges, block))
            inner = dict(layout, left=block.left, right=block.right, slack=0.0)
            # 每一格自己算和上一列的距離,才不會整個表格被墊高
            _write_block(cell, block, inner, {"baseline": previous})
        previous = max((b.lines[-1].baseline if b.kind == "text" else b.top + b.height) for b in row["blocks"])
    for row in table.rows:
        for cell in row.cells:
            # 儲存格一開始就有一個空段落:有內容就刪掉,沒內容的壓到 1 點,不然整列會被撐高
            if len(cell.paragraphs) > 1 and not cell.paragraphs[0].runs:
                element = cell.paragraphs[0]._element
                element.getparent().remove(element)
            elif not cell.paragraphs[0].runs:
                empty = cell.paragraphs[0].paragraph_format
                empty.space_before = empty.space_after = Pt(0)
                empty.line_spacing = Pt(1)
                cell.paragraphs[0].add_run().font.size = Pt(1)
    state["baseline"] = group[-1]["bottom"]




def _page_margins(section, width, blocks):
    if blocks:
        content_left = min(b.left for b in blocks)
        content_right = max(b.right for b in blocks)
        content_top = min(_block_span(b)[0] for b in blocks)
    else:
        content_left, content_right, content_top = MIN_MARGIN, width - MIN_MARGIN, MIN_MARGIN
    section.left_margin = Emu(round(max(MIN_MARGIN, content_left) * EMU_PER_POINT))
    section.right_margin = Emu(round(max(MIN_MARGIN, width - content_right) * EMU_PER_POINT))
    section.top_margin = Emu(round(max(MIN_MARGIN, content_top) * EMU_PER_POINT))
    section.bottom_margin = Emu(round(MIN_MARGIN * EMU_PER_POINT))
    return content_left, content_right


def _write_flow(word, width, section, blocks, rulings):
    """重新排版:輸出成一般的段落與表格,最好編輯。"""
    content_left, content_right = _page_margins(section, width, blocks)
    layout = {"left": content_left, "right": content_right, "pitch": _page_pitch(blocks),
              "slack": round(max(6.0, (content_right - content_left) * 0.05), 1)}
    state = {"baseline": None}
    previous_was_table = False
    for group in _group_tables(_group_rows(blocks)):
        if len(group) == 1 and len(group[0]["blocks"]) == 1:
            _write_block(word, group[0]["blocks"][0], layout, state)
            previous_was_table = False
        else:
            if previous_was_table:
                # 兩個表格直接相鄰時 Word 會把它們併成一個,中間放一個極小的空段落隔開
                spacer = word.add_paragraph().paragraph_format
                spacer.space_before = spacer.space_after = Pt(0)
                spacer.line_spacing = Pt(1)
            _write_table(word, group, layout, state, _has_rulings(group, rulings))
            previous_was_table = True


def _twips(points) -> str:
    return str(max(0, round(points * 20)))


def _frame(paragraph, left, top, width):
    """把段落固定在頁面上的位置(Word 的「框架」);不佔版面,也不會被其他內容推走。"""
    frame = OxmlElement("w:framePr")
    frame.set(qn("w:w"), _twips(width))
    frame.set(qn("w:x"), _twips(left))
    frame.set(qn("w:y"), _twips(top))
    frame.set(qn("w:hAnchor"), "page")
    frame.set(qn("w:vAnchor"), "page")
    frame.set(qn("w:wrap"), "around")
    paragraph._p.get_or_add_pPr().insert(0, frame)


def _write_exact(word, width, section, blocks):
    """照原樣:每一行字、每張圖都固定在原檔的位置;最像原檔,但改字時不會自動重排。

    每一行各自是一個框架:同一段的行左緣常常不一樣(縮排、項目符號),合成一個框架反而會對不齊。
    """
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Emu(round(MIN_MARGIN * EMU_PER_POINT)))
    page_pitch = _page_pitch(blocks)
    for block in sorted(blocks, key=lambda b: (_block_span(b)[0], b.left)):
        if block.kind == "image":
            paragraph = word.add_paragraph()
            spacing = paragraph.paragraph_format
            spacing.space_before = spacing.space_after = Pt(0)
            _frame(paragraph, block.left, block.top, block.width)
            paragraph.add_run().add_picture(io.BytesIO(block.image), width=Emu(round(block.width * EMU_PER_POINT)))
            continue
        pitch = _line_pitch(block, page_pitch)
        for line in block.lines:
            if not line.text.strip():
                continue
            paragraph = word.add_paragraph()
            spacing = paragraph.paragraph_format
            spacing.space_before = spacing.space_after = Pt(0)
            spacing.line_spacing = Pt(round(pitch, 1))
            # 固定行高時,字的底線大約落在這一行由上往下 80% 的地方
            _frame(paragraph, line.left, line.baseline - pitch * 0.8, max(line.right - line.left, width - line.left - 2))
            _write_runs(paragraph, Block("text", line.top, line.left, line.right, lines=[line]))
    # 每頁最後放一個不佔空間的空段落,換頁的設定才不會掛在框架上
    tail = word.add_paragraph().paragraph_format
    tail.space_before = tail.space_after = Pt(0)
    tail.line_spacing = Pt(1)


def _write_page(word, number, width, height, chars, images, scan, rulings=(), layout="flow", split_gap=2.5):
    section = word.sections[0] if number == 0 else word.add_section(WD_SECTION.NEW_PAGE)
    section.page_width, section.page_height = Emu(round(width * EMU_PER_POINT)), Emu(round(height * EMU_PER_POINT))

    if scan is not None:
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Emu(0)
        word.add_picture(io.BytesIO(scan), width=Emu(round(width * EMU_PER_POINT)))
        return

    blocks = _group_paragraphs(_group_lines(chars, split_gap)) + list(images)
    if layout == "exact":
        _write_exact(word, width, section, blocks)
    else:
        _write_flow(word, width, section, blocks, rulings)


# ---------------------------------------------------------------- 掃描檔(文字辨識)

def _is_scan(chars, images, width, height) -> bool:
    """整頁沒有文字,而且只有蓋住大半頁的圖(或什麼都沒有):多半是掃描檔。"""
    if chars:
        return False
    covered = sum(image.width * image.height for image in images)
    return not images or covered >= width * height * 0.6


def _ocr_size(line_words) -> float:
    """推估一行的字級(像素)。

    中文字是全形,相鄰兩個字左緣的距離正好是一個字寬,也就是字級,這個最準;
    沒有相鄰的中文字時才用字高推估(中文字的方框約是字級的 0.88 倍,英文約 0.72 倍)。
    """
    ordered = sorted(line_words, key=lambda w: w.left)
    advances = []
    for word, after in zip(ordered, ordered[1:]):
        if all(is_cjk(c) for c in word.text) and is_cjk(after.text[0]):
            step = (after.left - word.left) / len(word.text)
            if step <= (word.bottom - word.top) * 1.6:      # 中間隔一段空白的不算
                advances.append(step)
    if len(advances) >= 2:
        return statistics.median(advances)
    heights = [w.bottom - w.top for w in line_words if any(is_cjk(c) for c in w.text)]
    if heights:
        return statistics.median(heights) / 0.88
    return statistics.median(w.bottom - w.top for w in line_words) / 0.72


def _ocr_chars(words, scale) -> list:
    """把辨識出來的詞換成和 PDF 文字一樣的 Char,後面分行、分段、表格的流程就能共用。"""
    chars = []
    by_line = {}
    for word in words:
        by_line.setdefault(word.line, []).append(word)
    rows = []
    for line_words in by_line.values():
        # 辨識時左右並排、高低不同的字可能被湊成同一行,依上下範圍再拆開
        groups = []
        for word in sorted(line_words, key=lambda w: w.bottom - w.top, reverse=True):
            middle = (word.top + word.bottom) / 2
            group = next((g for g in groups if g["top"] <= middle <= g["bottom"]), None)
            if group is None:
                groups.append({"top": word.top, "bottom": word.bottom, "words": [word]})
            else:
                group["words"].append(word)
        rows += [(g["bottom"], g["words"]) for g in groups]
    for bottom, line_words in rows:
        size = max(4.0, round(_ocr_size(line_words) * scale * 2) / 2)
        font = CJK_FALLBACK["serif"]
        previous = None
        for word in sorted(line_words, key=lambda w: w.left):
            baseline = bottom * scale
            # 只在一般字距時補空白;空得很開的是表格的不同格,補了空白就分不開了
            if previous is not None and (word.left - previous.right) * scale < size * OCR_SPLIT_GAP:
                chars.append(Char(" ", previous.right * scale, word.top * scale, word.left * scale,
                                  word.bottom * scale, baseline, size, font, False, False, (0, 0, 0), True))
            step = (word.right - word.left) / len(word.text)
            for index, text in enumerate(word.text):
                left = (word.left + step * index) * scale
                chars.append(Char(text, left, word.top * scale, left + step * scale, word.bottom * scale,
                                  baseline, size, font, False, False, (0, 0, 0)))
            previous = word
    return chars


def _inside_any(char, figures) -> bool:
    """字有一半以上落在某張圖裡;圖裡的標籤已經在截圖上了,不要再變成段落。"""
    area = max(1e-6, (char.right - char.left) * (char.bottom - char.top))
    for f in figures:
        overlap_x = min(char.right, f.right) - max(char.left, f.left)
        overlap_y = min(char.bottom, f.top + f.height) - max(char.top, f.top)
        if overlap_x > 0 and overlap_y > 0 and overlap_x * overlap_y >= area * 0.5:
            return True
    return False


def _scan_figures(image, words, scale) -> list:
    """掃描頁上文字以外的圖(照片、圖表、印章):把文字塗白,剩下成塊的深色區域切出來當圖片。"""
    from PIL import Image, ImageDraw

    gray = image.convert("L")
    draw = ImageDraw.Draw(gray)
    for word in words:
        pad = max(2, (word.bottom - word.top) // 4)
        draw.rectangle((word.left - pad, word.top - pad, word.right + pad, word.bottom + pad), fill=255)
    cell = 8                                    # 縮小 8 倍再找,速度快,也順便忽略掃描的小雜點
    small = gray.resize((max(1, gray.width // cell), max(1, gray.height // cell)), Image.Resampling.BOX)
    w, h = small.size
    data = small.tobytes()
    # 比紙張顏色深一點就算:淺色底的區塊(投影片的圖框、色塊)連同裡面的標籤會整塊變成圖片,不會被拆散;
    # 紙色用出現最多的亮度,泛黃的掃描紙也不會整頁被當成圖
    histogram = small.histogram()
    paper = max(range(256), key=lambda value: histogram[value])
    ink = bytearray(1 if value < paper - 8 else 0 for value in data)
    boxes = []
    for start in range(w * h):
        if not ink[start]:
            continue
        ink[start] = 0
        stack = [start]
        left = right = start % w
        top = bottom = start // w
        count = 0
        while stack:
            at = stack.pop()
            cx, cy = at % w, at // w
            count += 1
            left, right, top, bottom = min(left, cx), max(right, cx), min(top, cy), max(bottom, cy)
            for ny in range(max(0, cy - 2), min(h, cy + 3)):     # 隔一兩格也算相連,圖裡的細線才不會斷開
                row = ny * w
                for nx in range(max(0, cx - 2), min(w, cx + 3)):
                    if ink[row + nx]:
                        ink[row + nx] = 0
                        stack.append(row + nx)
        boxes.append((left, top, right + 1, bottom + 1, count))
    figures = []
    min_side = 36 / scale / cell                # 小於半英吋的都當成雜點或殘留的筆畫
    # 被另一塊整個包住的(例如框裡的照片和外框之間隔著一圈白邊)已經在外框的截圖裡了
    boxes = [box for box in boxes
             if not any(other is not box and other[0] <= box[0] and other[1] <= box[1]
                        and other[2] >= box[2] and other[3] >= box[3] for other in boxes)]
    for left, top, right, bottom, count in boxes:
        if right - left < min_side or bottom - top < min_side or count < 12:
            continue
        if (right - left) * (bottom - top) > w * h * 0.7:
            continue                            # 幾乎整頁:多半是紙張底色不均,不是圖
        crop = image.crop((left * cell, top * cell, right * cell, bottom * cell)).convert("RGB")
        buffer = io.BytesIO()
        crop.save(buffer, "JPEG", quality=88)
        figures.append(Block("image", top * cell * scale, left * cell * scale, right * cell * scale,
                             image=buffer.getvalue(), width=(right - left) * cell * scale,
                             height=(bottom - top) * cell * scale))
    return figures


# ---------------------------------------------------------------- 主流程

def convert(source, out_dir, progress=None, cancel=None, layout="flow", ocr=None) -> Path:
    """把 PDF 轉成 Word;回傳產生的檔案。

    layout:"flow" 重新排版(好編輯)、"exact" 照原樣(每塊固定在原位)。
    ocr:文字辨識模組(plugins/documents/ocr.py);有給就會辨識掃描頁,沒給則把掃描頁整頁放成圖片。
    """
    source = Path(source)
    document = pdfium.open_document(source.read_bytes())
    word = Document()
    _prepare_document(word)
    try:
        total = pdfium.page_count(document)
        for number in range(total):
            if cancel is not None and cancel.is_set():
                raise InterruptedError("已取消")
            if progress:
                progress(number, total, f"第 {number + 1} 頁")
            scan = page_image = None
            split_gap = 2.5
            with pdfium.LOCK:
                page = document[number]
                try:
                    width, height = page.get_size()
                    text_page = page.get_textpage()
                    try:
                        chars = _read_chars(text_page.raw, height)
                    finally:
                        text_page.close()
                    images = _read_images(document, page, height)
                    rulings = _read_rulings(page, height)
                    if _is_scan(chars, images, width, height):
                        dpi = ocr.DPI if ocr is not None else SCAN_DPI
                        bitmap = page.render(scale=dpi / 72)
                        page_image = bitmap.to_pil().convert("RGB").copy()
                        bitmap.close()
                    else:
                        charts, chars = _vector_pictures(page, chars, width, height)
                        images += charts
                finally:
                    page.close()
            if page_image is not None:
                words = []
                if ocr is not None:
                    if progress:
                        progress(number, total, f"辨識第 {number + 1} 頁的文字")
                    words = ocr.recognize(page_image)
                if words:
                    scale = 72 / ocr.DPI
                    images = _scan_figures(page_image, words, scale)
                    chars = [c for c in _ocr_chars(words, scale) if not _inside_any(c, images)]
                    rulings = []
                    split_gap = OCR_SPLIT_GAP
                else:
                    # 沒有辨識(或辨識不到字):整頁照樣放成圖片
                    if ocr is not None:
                        page_image = page_image.resize((round(width * SCAN_DPI / 72), round(height * SCAN_DPI / 72)))
                    buffer = io.BytesIO()
                    page_image.save(buffer, "JPEG", quality=85)
                    scan = buffer.getvalue()
            _write_page(word, number, width, height, chars, images, scan, rulings, layout, split_gap)
    finally:
        pdfium.close(document)
    out = free_path(Path(out_dir), source.stem, ".docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    word.save(str(out))
    return out
